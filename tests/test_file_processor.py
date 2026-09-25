"""Regression tests: files handed to Gemini must arrive as real Parts; CSV/TSV must parse."""
import io
from types import SimpleNamespace

import pytest
from google.genai import types

from actions import file_processor as fp
from core import gemini


@pytest.fixture
def sent(monkeypatch):
    calls = []
    monkeypatch.setattr(fp.gemini, "call",
                        lambda contents, **k: calls.append(contents) or SimpleNamespace(text="ok"))
    return calls


def _media_parts(contents):
    return [p for p in gemini._to_live_parts(contents) if "inline_data" in p]


def test_image_reaches_the_model(tmp_path, sent):
    from PIL import Image
    path = tmp_path / "receipt.png"
    Image.new("RGB", (4, 4), "white").save(path)

    fp._process_image(path, "ocr", {"save": False})

    (contents,) = sent
    assert all(isinstance(c, (str, types.Part)) for c in contents)
    assert [p["inline_data"]["mime_type"] for p in _media_parts(contents)] == ["image/png"]


def test_audio_is_sent_as_a_valid_part(tmp_path, sent):
    path = tmp_path / "memo.mp3"
    path.write_bytes(b"ID3fake")

    fp._process_audio(path, "transcribe", {"save": False})

    (contents,) = sent
    assert all(isinstance(c, (str, types.Part)) for c in contents)
    assert [p["inline_data"]["mime_type"] for p in _media_parts(contents)] == ["audio/mpeg"]


@pytest.mark.parametrize("name,body", [("sales.csv", "city,total\nIzmir,3\nBursa,5\n"),
                                       ("sales.tsv", "city\ttotal\nIzmir\t3\nBursa\t5\n")])
def test_csv_and_tsv_parse_into_columns(tmp_path, name, body):
    pytest.importorskip("pandas")
    path = tmp_path / name
    path.write_text(body)

    out = fp._process_data(path, "csv", "info", {})

    assert "Could not read" not in out
    assert "total" in out and "2" in out


def test_free_form_action_is_sent_as_the_instruction(tmp_path, sent):
    path = tmp_path / "notes.txt"
    path.write_text("meeting at 10, lunch at 1")
    fp._process_text_doc(path, "text", "list the times", {"save": False})
    assert sent[0].startswith("list the times")


def test_unreadable_docx_is_reported_not_summarised(tmp_path, sent):
    path = tmp_path / "old.docx"
    path.write_bytes(b"not a zip")
    out = fp._process_text_doc(path, "docx", "summarize", {})
    assert "Could not read" in out and sent == []


def test_xml_is_processed_as_text(tmp_path, sent):
    path = tmp_path / "feed.xml"
    path.write_text("<rss><item>hi</item></rss>")
    out = fp.file_processor({"file_path": str(path), "action": "word_count"})
    assert "Invalid JSON" not in out and "Word count" in out


def test_filter_equals_matches_numbers(tmp_path):
    pytest.importorskip("pandas")
    path = tmp_path / "t.csv"
    path.write_text("id,total\n1,5\n2,7\n")
    assert "1 rows match" in fp._process_data(path, "csv", "filter",
                                               {"column": "total", "value": "5"})


def test_sorting_excel_writes_a_real_workbook(tmp_path):
    pd = pytest.importorskip("pandas")
    path = tmp_path / "t.xlsx"
    pd.DataFrame({"n": [3, 1, 2]}).to_excel(path, index=False)
    fp._process_data(path, "excel", "sort", {"column": "n"})
    assert list(pd.read_excel(tmp_path / "t_sorted.xlsx")["n"]) == [1, 2, 3]


def test_failed_ffmpeg_run_is_not_reported_as_success(tmp_path, monkeypatch):
    import subprocess
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"x")
    def fake_run(argv, **k):
        if "-version" in argv:
            return SimpleNamespace(returncode=0)            # ffmpeg is installed
        if k.get("check"):
            raise subprocess.CalledProcessError(1, argv)
        return SimpleNamespace(returncode=1)                # ...but this run fails
    monkeypatch.setattr(fp.subprocess, "run", fake_run)
    assert "failed" in fp._process_video(video, "convert", {"format": "webm"}).lower()


def test_png_with_alpha_converts_to_jpeg(tmp_path):
    from PIL import Image
    src = tmp_path / "logo.png"
    Image.new("RGBA", (4, 4), (255, 0, 0, 128)).save(src)
    assert "Converted" in fp._process_image(src, "convert", {"format": "jpeg"})


def test_numeric_headers_and_unknown_formats_do_not_crash(tmp_path):
    pd = pytest.importorskip("pandas")
    path = tmp_path / "t.xlsx"
    pd.DataFrame({1: [1], 2: [2]}).to_excel(path, index=False)
    assert "Available" in fp._process_data(path, "excel", "filter", {"column": "x"})
    assert "Unsupported format" in fp._process_data(path, "excel", "convert", {"format": "pdf"})


def test_time_strings_parse_to_seconds():
    assert fp._seconds("00:01:30") == 90 and fp._seconds("1:30") == 90 and fp._seconds(90) == 90


def test_json_with_bom_is_valid(tmp_path, sent):
    path = tmp_path / "d.json"
    path.write_bytes(b"\xef\xbb\xbf" + b'{"a": 1}')
    assert "Invalid JSON" not in fp._process_json(path, "analyze", {})


def test_extract_text_of_a_text_file_returns_it(tmp_path):
    path = tmp_path / "n.txt"
    path.write_text("hello there")
    assert fp._process_text_doc(path, "text", "extract_text", {}) == "hello there"
    assert not list(tmp_path.glob("*_extracted.txt"))


def test_docx_tables_are_read(tmp_path):
    docx = pytest.importorskip("docx")
    doc = docx.Document()
    doc.add_paragraph("Invoice 42")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Total"
    t.cell(1, 0).text, t.cell(1, 1).text = "Consulting", "€1,250"
    path = tmp_path / "invoice.docx"
    doc.save(path)
    fp._process_text_doc(path, "docx", "extract_text", {})
    out = (tmp_path / "invoice_extracted.txt").read_text(encoding="utf-8")
    assert "Invoice 42" in out and "Consulting | €1,250" in out


def test_pptx_tables_and_groups_are_read(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(1, 0).text, table.cell(1, 1).text = "Q3 revenue", "4.2M"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = "Grouped note"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    fp._process_pptx(path, "extract_text", {})
    out = (tmp_path / "deck_text.txt").read_text(encoding="utf-8")
    assert "Q3 revenue | 4.2M" in out and "Grouped note" in out
